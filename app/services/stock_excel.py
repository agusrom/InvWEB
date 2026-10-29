from dataclasses import dataclass, field
from io import BytesIO
import math
import re

from openpyxl import load_workbook
import xlrd


ENCABEZADOS = {
    "PT": {"gcod", "ccod", "stok", "nueva_cant"},
    "PRD": {
        "piezacod",
        "codigocliente",
        "operacion",
        "proceso",
        "cantidad",
        "retorno",
        "proveedor",
    },
    "MP": {"lote", "cmat", "descr", "undad", "cant"},
    "ASIGNADO": {"ccodmat", "ncantidad"},
}


@dataclass
class StockPRD:
    cantidad: float = 0
    proveedores: set[str] = field(default_factory=set)
    procesos: set[str] = field(default_factory=set)


@dataclass
class StockMPLote:
    cantidad: float = 0
    descripcion: str = ""
    unidad: str = ""


@dataclass
class ReporteStockExcel:
    tipo: str
    hoja: str
    filas: int
    pt: dict[str, float] = field(default_factory=dict)
    prd: dict[str, dict[str, StockPRD]] = field(default_factory=dict)
    mp: dict[str, dict[str, StockMPLote]] = field(default_factory=dict)
    asignado: dict[str, float] = field(default_factory=dict)


def texto_codigo(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and math.isnan(valor):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor)).strip().upper()
    return str(valor).strip().upper()


def texto_proveedor(valor) -> str:
    texto = texto_codigo(valor)
    return "" if texto in {"", "0", "0.0"} else texto


def texto_lote(valor) -> str:
    lote = texto_codigo(valor)
    if lote.isdigit() and len(lote) < 5:
        lote = lote.zfill(5)
    return lote


def numero(valor) -> float:
    if valor in (None, ""):
        return 0.0
    if isinstance(valor, (int, float)):
        if isinstance(valor, float) and math.isnan(valor):
            return 0.0
        return float(valor)
    texto = str(valor).strip().replace(" ", "")
    if not texto:
        return 0.0
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError as error:
        raise ValueError(f"Cantidad no numérica: {valor}") from error


def _leer_filas_excel(contenido: bytes, tipo: str):
    """Devuelve nombre de hoja y filas para reportes .xls o .xlsx."""
    try:
        libro = xlrd.open_workbook(file_contents=contenido)
        if not libro.sheet_names():
            raise ValueError(f"El archivo {tipo} no tiene hojas")
        hoja = libro.sheet_by_index(0)
        return hoja.name, [
            [hoja.cell_value(fila, columna) for columna in range(hoja.ncols)]
            for fila in range(hoja.nrows)
        ]
    except ValueError:
        raise
    except Exception:
        try:
            libro = load_workbook(
                BytesIO(contenido),
                read_only=True,
                data_only=True,
            )
            hoja = libro.worksheets[0]
            return hoja.title, [list(fila) for fila in hoja.iter_rows(values_only=True)]
        except Exception as error:
            raise ValueError(
                f"El archivo {tipo} no es un Excel .xls o .xlsx válido"
            ) from error


def leer_reporte_stock(contenido: bytes, tipo: str) -> ReporteStockExcel:
    tipo = tipo.upper()
    if tipo not in ENCABEZADOS:
        raise ValueError("Tipo de reporte no válido")
    if not contenido:
        raise ValueError(f"El archivo {tipo} está vacío")

    nombre_hoja, filas = _leer_filas_excel(contenido, tipo)
    if not filas:
        raise ValueError(f"El archivo {tipo} no tiene encabezados")
    encabezados = {
        re.sub(r"\s+", "", str(valor or "").strip().lower()):
        columna
        for columna, valor in enumerate(filas[0])
        if str(valor or "").strip()
    }
    faltantes = ENCABEZADOS[tipo] - set(encabezados)
    if faltantes:
        raise ValueError(
            f"El reporte {tipo} no tiene las columnas: "
            + ", ".join(sorted(faltantes))
        )

    reporte = ReporteStockExcel(
        tipo=tipo,
        hoja=nombre_hoja,
        filas=max(0, len(filas) - 1),
    )
    for fila in filas[1:]:
        valor = lambda nombre: (
            fila[encabezados[nombre]]
            if encabezados[nombre] < len(fila)
            else None
        )
        if tipo == "PT":
            codigo = texto_codigo(valor("gcod"))
            if codigo:
                reporte.pt[codigo] = (
                    reporte.pt.get(codigo, 0) + numero(valor("stok"))
                )
        elif tipo == "PRD":
            codigo = texto_codigo(valor("piezacod"))
            operacion = texto_codigo(valor("operacion"))
            if not codigo or not operacion:
                continue
            dato = reporte.prd.setdefault(codigo, {}).setdefault(
                operacion,
                StockPRD(),
            )
            dato.cantidad += numero(valor("cantidad"))
            proveedor = texto_proveedor(valor("proveedor"))
            proceso = texto_codigo(valor("proceso"))
            if proveedor:
                dato.proveedores.add(proveedor)
            if proceso:
                dato.procesos.add(proceso)
        elif tipo == "MP":
            codigo = texto_codigo(valor("cmat"))
            lote = texto_lote(valor("lote"))
            if not codigo:
                continue
            if not re.fullmatch(r"\d{5}", lote):
                raise ValueError(
                    f"Lote inválido para {codigo}: {lote or 'vacío'}. "
                    "Debe tener 5 dígitos."
                )
            dato = reporte.mp.setdefault(codigo, {}).setdefault(
                lote,
                StockMPLote(),
            )
            dato.cantidad += numero(valor("cant"))
            descripcion = str(valor("descr") or "").strip()
            unidad = texto_codigo(valor("undad"))
            if descripcion:
                dato.descripcion = descripcion
            if unidad:
                dato.unidad = unidad
        else:
            codigo = texto_codigo(valor("ccodmat"))
            if codigo:
                reporte.asignado[codigo] = (
                    reporte.asignado.get(codigo, 0)
                    + numero(valor("ncantidad"))
                )
    return reporte


def proveedores_texto(dato: StockPRD) -> str:
    return ", ".join(sorted(dato.proveedores))
