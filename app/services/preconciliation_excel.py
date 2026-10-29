"""Exportaciones de las dos vistas del módulo Análisis de stock."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.preconciliation_indicators import calcular_indicadores_preconciliacion

BLUE = "244F7E"
LIGHT_BLUE = "EAF0F8"
GREEN = "E6F6F2"
RED = "FFF0F0"
GRAY = "F2F4F7"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D0D5DD")


def _title(sheet, title, subtitle, columns):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    cell = sheet.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="Arial", size=16, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    sheet.cell(2, 1, subtitle).font = Font(name="Arial", size=10, bold=True, color=BLUE)


def _headers(sheet, row, headers):
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(row, column, label)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(name="Arial", size=9, bold=True, color=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    sheet.row_dimensions[row].height = 30
    sheet.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"
    sheet.freeze_panes = f"C{row + 1}"


def _finish(sheet, number_columns=()):
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
        width = min(max(max((len(value) for value in values), default=8) + 2, 10), 34)
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in number_columns:
        for row in range(4, sheet.max_row + 1):
            sheet.cell(row, column).number_format = "#,##0.00"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _diagnosis(rows):
    system_ok = bool(rows) and rows[0].get("sistema_ok", False)
    reconciliations = [row.get("conciliacion_ok", False) for row in rows]
    has_adjustments = any(row.get("ajustes_sistema") for row in rows)
    if system_ok and reconciliations and all(reconciliations):
        return "Ambos controles correctos"
    if not system_ok and reconciliations and all(reconciliations) and has_adjustments:
        return "Revisar ajustes del sistema"
    if not system_ok and reconciliations and all(reconciliations):
        return "Revisar gestión del sistema"
    if system_ok and not all(reconciliations):
        return "Revisar documentación"
    if any(reconciliations):
        return "Desvío localizado"
    return "Revisión completa"


def _analysis_sheet(book, final, calculation, notes, history):
    sheet = book.active
    sheet.title = "Análisis general"
    scenarios = calculation["escenarios"]
    scenario_maps = [{row["codigo"]: row for row in scenario.get("filas", [])} for scenario in scenarios]
    codes = sorted(set().union(*(set(items) for items in scenario_maps))) if scenario_maps else []
    headers = ["Código interno", "Código cliente", "Descripción", "Diagnóstico", "Sistema directo comparable", "Sistema MP cantidad", "CONS_", "UMCO_", "Sistema MP equivalente", "Sistema en proceso", "Sistema PT", "Stock asignado", "Físico directo comparable", "Físico MP cantidad", "Físico MP equivalente", "Físico en proceso", "Físico PT", "Dif. sistema directo", "Conteo actual explotado"]
    for item in history:
        headers += [f"{item['codigo']} conteo", f"{item['codigo']} sistema", f"{item['codigo']} diferencia"]
    for scenario in scenarios:
        prefix = scenario["inicial"]["codigo"]
        headers += [f"{prefix} esperado explotado", f"{prefix} físico explotado vs esperado"]
    headers += ["Estado", "Causa probable", "Observación"]
    _title(sheet, f"ANÁLISIS GENERAL · {final.codigo}", f"{final.cliente} · {final.fecha.strftime('%d/%m/%Y')}", len(headers))
    _headers(sheet, 4, headers)
    for row_number, code in enumerate(codes, 5):
        rows = [items[code] for items in scenario_maps if code in items]
        current = rows[0]
        detail = current.get("detalle_sistema", {})
        physical = current.get("detalle_fisico", {})
        values = [code, current.get("codigo_cliente", ""), current.get("descripcion", ""), _diagnosis(rows), current.get("sistema_actual", 0), detail.get("materia_prima", 0), detail.get("consumo_material", 1), detail.get("unidad_consumo", "C/U"), detail.get("materia_prima_equivalente", 0), detail.get("en_proceso", 0), detail.get("producto_terminado", 0), detail.get("stock_asignado", 0), current.get("conteo_actual", 0), physical.get("materia_prima", 0), physical.get("materia_prima_equivalente", 0), physical.get("en_proceso", 0), physical.get("producto_terminado", 0), None, current.get("conteo_explotado_actual", current.get("fisico_actual", 0))]
        for historic in history:
            item = historic.get("filas", {}).get(code)
            values += [item.get("conteo") if item else None, item.get("sistema") if item else None, None]
        for scenario_map in scenario_maps:
            item = scenario_map.get(code)
            values += [item.get("esperado") if item else None, None]
        note = notes.get(code, {})
        values += [note.get("estado", "PENDIENTE"), note.get("causa", ""), note.get("observacion", "")]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
        # Diferencias quedan auditables como fórmulas.
        sheet.cell(row_number, 18, f"=M{row_number}-E{row_number}")
        column = 20
        for _ in history:
            sheet.cell(row_number, column + 2, f"={get_column_letter(column)}{row_number}-{get_column_letter(column + 1)}{row_number}")
            column += 3
        for _ in scenarios:
            expected_column = get_column_letter(column)
            sheet.cell(row_number, column + 1, f"=S{row_number}-{expected_column}{row_number}")
            column += 2
        if current.get("diferencia_sistema"):
            sheet.cell(row_number, 18).fill = PatternFill("solid", fgColor=RED)
    _finish(sheet, range(5, len(headers) - 2))


def _reconciliation_summary(book, final, scenario, notes):
    sheet = book.active
    sheet.title = "Conciliación"
    headers = ["Código interno", "Código cliente", "Descripción", "Inicial explotado", "Ingresos", "Egresos", "Scrap", "Esperado explotado", "Conteo explotado actual", "Físico explotado vs esperado", "Estado", "Causa", "Observación"]
    _title(sheet, f"CONCILIACIÓN · {scenario['inicial']['codigo']} → {final.codigo}", f"{scenario['inicial']['fecha']} al {final.fecha.strftime('%d/%m/%Y')}", len(headers))
    _headers(sheet, 4, headers)
    for row_number, item in enumerate(scenario.get("filas", []), 5):
        note = notes.get(item["codigo"], {})
        values = [item["codigo"], item.get("codigo_cliente", ""), item.get("descripcion", ""), item.get("fisico_inicial", 0), item.get("ingresos", 0), item.get("egresos", 0), item.get("scrap", 0), None, item.get("conteo_explotado_actual", item.get("fisico_actual", 0)), None, note.get("estado", "PENDIENTE"), note.get("causa", ""), note.get("observacion", "")]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
        sheet.cell(row_number, 8, f"=D{row_number}+E{row_number}-F{row_number}-G{row_number}")
        sheet.cell(row_number, 10, f"=I{row_number}-H{row_number}")
        if item.get("diferencia_conciliacion"):
            sheet.cell(row_number, 10).fill = PatternFill("solid", fgColor=RED)
        else:
            sheet.cell(row_number, 10).fill = PatternFill("solid", fgColor=GREEN)
    _finish(sheet, range(4, 11))


def _detail_sheet(book, name, scenario, kind):
    sheet = book.create_sheet(name)
    if kind == "IN":
        headers = ["Código interno destino", "Código cliente destino", "Descripción", "Fecha", "Remito", "Material", "Código cliente origen", "Cantidad", "Verificado", "Observación"]
    else:
        headers = ["Código interno destino", "Código cliente destino", "Descripción", "Fecha", "Remito", "Pieza origen", "Código cliente origen", "Cantidad pieza", "Uso", "Cantidad equivalente", "Verificado", "Observación"]
    _headers(sheet, 1, headers)
    output_row = 2
    for item in scenario.get("filas", []):
        details = item.get("detalle_ingresos" if kind == "IN" else "detalle_egresos", [])
        for detail in details:
            if kind == "IN":
                values = [item["codigo"], item.get("codigo_cliente", ""), item.get("descripcion", ""), detail.get("fecha"), detail.get("remito"), detail.get("codigo_material", detail.get("codigo")), detail.get("codigo_cliente_origen", ""), detail.get("cantidad", 0), detail.get("verificado", False), detail.get("observacion", "")]
            else:
                values = [item["codigo"], item.get("codigo_cliente", ""), item.get("descripcion", ""), detail.get("fecha"), detail.get("remito"), detail.get("pieza_origen"), detail.get("codigo_cliente_origen", ""), detail.get("cantidad_pieza", 0), detail.get("uso", 0), detail.get("cantidad_equivalente", 0), detail.get("verificado", False), detail.get("observacion", "")]
            for column, value in enumerate(values, 1):
                sheet.cell(output_row, column, value)
            output_row += 1
    _finish(sheet, (8,) if kind == "IN" else (8, 9, 10))


def _audit_sheet(book, movements):
    sheet = book.create_sheet("Auditoría remitos")
    headers = ["Tipo", "Fecha", "Remito", "Código interno", "Código cliente", "Cantidad original", "Cantidad efectiva", "Manual", "Verificado", "Observación", "Escenarios incluidos", "Escenarios excluidos"]
    _headers(sheet, 1, headers)
    for row_number, item in enumerate(movements, 2):
        values = ["Ingreso" if item["tipo"] == "IN" else "Egreso", item.get("fecha"), item.get("remito"), item.get("codigo"), item.get("codigo_cliente", ""), item.get("cantidad_original", 0), item.get("cantidad", 0), item.get("manual", False), item.get("verificado", False), item.get("observacion", ""), ", ".join(map(str, item.get("escenarios", []))), ", ".join(map(str, item.get("excluidos", [])))]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
    _finish(sheet, (6, 7))


def _adjustment_sheet(book, scenario):
    sheet = book.create_sheet("Ajustes sistema")
    headers = ["Código interno", "Código cliente", "Descripción", "Fecha", "Operación", "Etapa", "Cantidad", "Motivo", "Afecta conciliación física"]
    _headers(sheet, 1, headers)
    output_row = 2
    for item in scenario.get("filas", []):
        for adjustment in item.get("ajustes_sistema", []):
            values = [item["codigo"], item.get("codigo_cliente", ""), item.get("descripcion", ""), adjustment.get("fecha"), adjustment.get("operacion"), "Producto terminado" if adjustment.get("etapa") == "TERMINADO" else "En proceso", adjustment.get("cantidad", 0), adjustment.get("motivo", ""), "No; se usa como evidencia"]
            for column, value in enumerate(values, 1):
                sheet.cell(output_row, column, value)
            output_row += 1
    _finish(sheet, (7,))


def _indicator_sheet(book, calculation, notes):
    sheet = book.create_sheet("Indicadores")
    indicators = calcular_indicadores_preconciliacion(calculation, notes)
    headers = ["Indicador", "%", "Buenas / completas", "Malas / pendientes", "Total", "Lectura"]
    _title(
        sheet,
        "INDICADORES · DOS CONTROLES INDEPENDIENTES",
        f"Sistema directo y conciliación explotada · pieza buena dentro de ±{indicators.get('tolerancia', 5)}%",
        len(headers),
    )
    _headers(sheet, 4, headers)
    metrics = [
        ("Confiabilidad vs. sistema", indicators["sistema"], "Conteo físico comparado con el stock del sistema"),
        ("Conciliación integral", indicators["conciliacion_integral"], "Piezas que cierran en todos los escenarios elegidos"),
        ("Aprueban ambos controles", indicators["consistencia_integral"], "Cumplen por separado sistema directo y conciliación explotada"),
        ("Remitos verificados", indicators["auditoria_documental"], "Movimientos de ingreso y egreso auditados"),
        ("Cobertura de causas", indicators["cobertura_causas"], "Piezas con diferencia y causa asignada"),
        ("Análisis gestionado", indicators["avance_analisis"], "Piezas con diferencia que ya salieron de Pendiente"),
        ("Casos resueltos", indicators["resolucion"], "Piezas con diferencia marcadas como resueltas"),
        ("Diferencias con ajustes", indicators["diferencias_con_ajustes"], "Diferencias directas con registros en produc16; no se comparan con el esperado"),
    ]
    for row_number, (name, metric, reading) in enumerate(metrics, 5):
        values = [name, None if metric["porcentaje"] is None else metric["porcentaje"] / 100, metric["buenas"], metric["malas"], metric["total"], reading]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
        sheet.cell(row_number, 2).number_format = "0.0%"
        if metric["porcentaje"] is not None and name != "Diferencias con ajustes":
            fill = GREEN if metric["porcentaje"] >= 95 else ("FFF4D8" if metric["porcentaje"] >= 80 else RED)
            sheet.cell(row_number, 2).fill = PatternFill("solid", fgColor=fill)
        elif name == "Diferencias con ajustes":
            sheet.cell(row_number, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    start = 15
    sheet.cell(start, 1, "CONFIABILIDAD POR INVENTARIO INICIAL").font = Font(bold=True, color=BLUE)
    scenario_headers = ["Inventario inicial", "Fecha", "% buenas", "Buenas", "Con desvío", "Total"]
    _headers(sheet, start + 1, scenario_headers)
    worst = indicators.get("peor_escenario") or {}
    for row_number, scenario in enumerate(indicators.get("escenarios", []), start + 2):
        values = [scenario.get("codigo"), scenario.get("fecha"), None if scenario.get("porcentaje") is None else scenario["porcentaje"] / 100, scenario.get("buenas", 0), scenario.get("malas", 0), scenario.get("total", 0)]
        for column, value in enumerate(values, 1):
            sheet.cell(row_number, column, value)
        sheet.cell(row_number, 3).number_format = "0.0%"
        if scenario.get("id") == worst.get("id"):
            for cell in sheet[row_number]:
                cell.fill = PatternFill("solid", fgColor="FFF4D8")
    _finish(sheet, (4, 5, 6))


def crear_excel_preconciliacion(final, calculation, notes, history, view):
    book = Workbook()
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    if view == "ANALISIS":
        _analysis_sheet(book, final, calculation, notes, history)
    else:
        scenario = calculation["escenarios"][0]
        _reconciliation_summary(book, final, scenario, notes)
        _detail_sheet(book, "Ingresos", scenario, "IN")
        _detail_sheet(book, "Egresos", scenario, "OUT")
        _audit_sheet(book, calculation.get("movimientos", []))
        _adjustment_sheet(book, scenario)
    _indicator_sheet(book, calculation, notes)
    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream
