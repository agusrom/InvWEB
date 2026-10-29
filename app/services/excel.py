from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from app.models import Inventario
from app.services.count_locations import desglosar_conteo_por_ubicacion
from app.serializers import (
    metricas_confiabilidad,
    metricas_conteo_item,
)


AZUL = "244F7E"
AZUL_OSCURO = "193754"
AZUL_CLARO = "EAF0F8"
AMARILLO = "FFC400"
GRIS = "F2F4F7"
GRIS_TEXTO = "667085"
VERDE = "E6F6F2"
VERDE_TEXTO = "08776D"
ROJO = "FFF0F0"
ROJO_TEXTO = "B42318"
NARANJA = "FFF4E5"
NARANJA_TEXTO = "9A5D00"
BLANCO = "FFFFFF"

BORDE_FINO = Side(style="thin", color="D0D5DD")
BORDE_AZUL = Side(style="thin", color="315D8B")

NOMBRES_ETAPA = {
    "MATERIAL": "Material sin usar",
    "PROCESO_GLOBAL": "Proceso global",
    "OPERACION": "En proceso",
    "TERMINADO": "Pieza terminada",
}


def texto_ambito(inventario: Inventario) -> str:
    return "Con cliente" if inventario.ambito == "CLIENTE" else "Interno"


def texto_etapa(etapa: str) -> str:
    return NOMBRES_ETAPA.get(etapa, etapa or "")


def texto_operacion(item) -> str:
    if item.etapa in {"PROCESO_GLOBAL", "TERMINADO"}:
        return ""
    if item.etapa == "MATERIAL":
        if item.lote:
            return f"Lote {item.lote}"
        titulo = item.titulo_operacion or ""
        if "Lotes:" in titulo:
            return f"Lotes: {titulo.split('Lotes:', 1)[1].strip()}"
        return ""
    return " - ".join(
        parte
        for parte in [
            (item.operacion or "").strip(),
            (item.titulo_operacion or "").strip(),
        ]
        if parte
    )


def codigo_resumen(item) -> str:
    if item.etapa == "MATERIAL" and item.codigo_material:
        return item.codigo_material
    return item.codigo


def descripcion_resumen(item) -> str:
    if item.etapa == "MATERIAL":
        return f"Materia prima de {item.codigo} - {item.descripcion}"
    return item.descripcion


def resultado_item(item, metricas: dict) -> str:
    if item.stock_pendiente:
        return "Stock pendiente"
    if item.conteo_fisico is None:
        return "Sin contar"
    if metricas["dentro_tolerancia"]:
        return "Dentro de ±5%"
    return "Fuera de ±5%"


def cantidad_columnas_resumen(inventario: Inventario) -> int:
    return 16 if inventario.modalidad_conteo == "MEDIOS" else 18


def aplicar_titulo(hoja, inventario: Inventario) -> None:
    ultima_columna = get_column_letter(cantidad_columnas_resumen(inventario))
    hoja.merge_cells(f"A1:{ultima_columna}1")
    hoja["A1"] = f"INVENTARIO {inventario.codigo}"
    hoja["A1"].fill = PatternFill("solid", fgColor=AZUL)
    hoja["A1"].font = Font(
        name="Arial",
        size=18,
        bold=True,
        color=BLANCO,
    )
    hoja["A1"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )
    hoja.row_dimensions[1].height = 32

    hoja.merge_cells(f"A2:{ultima_columna}2")
    hoja["A2"] = (
        f"{inventario.cliente}  |  "
        f"Realizado {inventario.fecha.strftime('%d/%m/%Y')}  |  "
        f"{texto_ambito(inventario)}"
    )
    hoja["A2"].font = Font(
        name="Arial",
        size=10,
        bold=True,
        color=AZUL_OSCURO,
    )
    hoja["A2"].alignment = Alignment(vertical="center")
    hoja.row_dimensions[2].height = 20


def escribir_campo(
    hoja,
    fila: int,
    columna: int,
    etiqueta: str,
    valor,
    ancho_valor: int = 2,
) -> None:
    celda_etiqueta = hoja.cell(fila, columna, etiqueta)
    celda_etiqueta.fill = PatternFill("solid", fgColor=AZUL_CLARO)
    celda_etiqueta.font = Font(
        name="Arial",
        size=8,
        bold=True,
        color=AZUL_OSCURO,
    )
    celda_etiqueta.alignment = Alignment(
        vertical="center",
        wrap_text=True,
    )
    celda_etiqueta.border = Border(
        left=BORDE_FINO,
        top=BORDE_FINO,
        bottom=BORDE_FINO,
    )

    inicio = columna + 1
    fin = columna + ancho_valor
    if fin > inicio:
        hoja.merge_cells(
            start_row=fila,
            start_column=inicio,
            end_row=fila,
            end_column=fin,
        )
    celda_valor = hoja.cell(fila, inicio, valor)
    celda_valor.font = Font(name="Arial", size=9, color="101828")
    celda_valor.alignment = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )
    celda_valor.border = Border(
        right=BORDE_FINO,
        top=BORDE_FINO,
        bottom=BORDE_FINO,
    )


def escribir_datos_generales(hoja, inventario: Inventario) -> None:
    cantidad_columnas = cantidad_columnas_resumen(inventario)
    ultima_columna = get_column_letter(cantidad_columnas)
    ancho_final = cantidad_columnas - 10
    escribir_campo(hoja, 4, 1, "Inventario", inventario.codigo, 2)
    escribir_campo(hoja, 4, 4, "Estado", inventario.estado, 2)
    escribir_campo(
        hoja,
        4,
        7,
        "Origen",
        getattr(inventario, "origen", "PLANTILLA"),
        2,
    )
    escribir_campo(
        hoja,
        4,
        10,
        "Categoría / modalidad",
        (
            f"{inventario.categoria} · MEDIOS"
            if inventario.modalidad_conteo == "MEDIOS"
            else inventario.categoria
        ),
        ancho_final,
    )

    escribir_campo(hoja, 5, 1, "Fecha realizada", inventario.fecha, 2)
    hoja["B5"].number_format = "dd/mm/yyyy"
    escribir_campo(
        hoja,
        5,
        4,
        "Fecha programada",
        inventario.fecha_programada or "No programado",
        2,
    )
    if inventario.fecha_programada:
        hoja["E5"].number_format = "dd/mm/yyyy"
    escribir_campo(
        hoja,
        5,
        7,
        "Cerrado el",
        inventario.cerrado_en or "Pendiente",
        2,
    )
    if inventario.cerrado_en:
        hoja["H5"].number_format = "dd/mm/yyyy hh:mm"
    escribir_campo(
        hoja,
        5,
        10,
        "Ámbito / mes de cierre",
        (
            f"{texto_ambito(inventario)} · "
            f"{(inventario.mes_cierre or inventario.fecha).strftime('%m/%Y')}"
        ),
        ancho_final,
    )

    escribir_campo(hoja, 6, 1, "Responsable", inventario.responsable, 2)
    escribir_campo(
        hoja,
        6,
        4,
        "Colaboradores",
        inventario.colaboradores or "Sin colaboradores",
        5,
    )
    escribir_campo(
        hoja,
        6,
        10,
        "Fuente de stock",
        inventario.stock_origen,
        ancho_final,
    )

    escribir_campo(
        hoja,
        7,
        1,
        "Plantilla ID",
        inventario.plantilla_id or "Sin plantilla",
        2,
    )
    escribir_campo(
        hoja,
        7,
        4,
        "Programación ID",
        inventario.programacion_id or "No programado",
        2,
    )
    escribir_campo(
        hoja,
        7,
        7,
        "Stock actualizado",
        inventario.stock_actualizado_en or "Sin registro",
        2,
    )
    if inventario.stock_actualizado_en:
        hoja["H7"].number_format = "dd/mm/yyyy hh:mm"
    escribir_campo(
        hoja,
        7,
        10,
        "Controla remitos",
        "Sí" if inventario.controla_remitos else "No",
        ancho_final,
    )

    if inventario.controla_remitos:
        escribir_campo(
            hoja,
            8,
            1,
            "Último remito ingreso",
            inventario.ultimo_remito_ingreso or "Sin informar",
            5,
        )
        escribir_campo(
            hoja,
            8,
            7,
            "Último remito egreso",
            inventario.ultimo_remito_egreso or "Sin informar",
            cantidad_columnas - 7,
        )
    else:
        hoja.merge_cells(f"A8:{ultima_columna}8")
        hoja["A8"] = "Este inventario no controla remitos."
        hoja["A8"].fill = PatternFill("solid", fgColor=GRIS)
        hoja["A8"].font = Font(
            name="Arial",
            size=8,
            italic=True,
            color=GRIS_TEXTO,
        )
        hoja["A8"].alignment = Alignment(vertical="center")

    hoja.merge_cells(f"A9:{ultima_columna}10")
    hoja["A9"] = (
        f"COMENTARIO\n{inventario.observaciones or 'Sin comentarios'}"
    )
    hoja["A9"].fill = PatternFill("solid", fgColor="F8FAFC")
    hoja["A9"].font = Font(name="Arial", size=9, color="344054")
    hoja["A9"].alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )
    hoja["A9"].border = Border(
        left=BORDE_FINO,
        right=BORDE_FINO,
        top=BORDE_FINO,
        bottom=BORDE_FINO,
    )
    hoja.row_dimensions[9].height = 22
    hoja.row_dimensions[10].height = 10


def escribir_indicadores(hoja, inventario: Inventario) -> None:
    confianza = metricas_confiabilidad(inventario)
    contados = sum(
        item.conteo_fisico is not None for item in inventario.items
    )
    diferencias = sum(
        metricas_conteo_item(item)["diferencia"] not in {None, 0}
        for item in inventario.items
    )
    total = len(inventario.items)
    indicadores = [
        ("PROGRESO", f"{contados} / {total}", AZUL_CLARO, AZUL_OSCURO),
        ("DIFERENCIAS", diferencias, NARANJA, NARANJA_TEXTO),
        (
            "CONFIABILIDAD (±5%)",
            (
                confianza["confiabilidad"] / 100
                if confianza["confiabilidad"] is not None
                else None
            ),
            VERDE if confianza["confiabilidad"] == 100 else ROJO,
            VERDE_TEXTO if confianza["confiabilidad"] == 100 else ROJO_TEXTO,
        ),
        (
            "RENGLONES OK",
            f'{confianza["items_ok"]} / {confianza["items_evaluables"]}',
            GRIS,
            AZUL_OSCURO,
        ),
    ]
    bloques = [
        (1, 3),
        (4, 6),
        (7, 9),
        (10, cantidad_columnas_resumen(inventario)),
    ]
    for (etiqueta, valor, fondo, color), (inicio, fin) in zip(
        indicadores,
        bloques,
    ):
        hoja.merge_cells(
            start_row=12,
            start_column=inicio,
            end_row=12,
            end_column=fin,
        )
        hoja.merge_cells(
            start_row=13,
            start_column=inicio,
            end_row=14,
            end_column=fin,
        )
        titulo = hoja.cell(12, inicio, etiqueta)
        titulo.fill = PatternFill("solid", fgColor=fondo)
        titulo.font = Font(
            name="Arial",
            size=8,
            bold=True,
            color=color,
        )
        titulo.alignment = Alignment(horizontal="center", vertical="center")
        valor_celda = hoja.cell(13, inicio, valor)
        valor_celda.fill = PatternFill("solid", fgColor=fondo)
        valor_celda.font = Font(
            name="Arial",
            size=16,
            bold=True,
            color=color,
        )
        valor_celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        if etiqueta.startswith("CONFIABILIDAD"):
            valor_celda.number_format = "0.00%"
        for fila in range(12, 15):
            for columna in range(inicio, fin + 1):
                hoja.cell(fila, columna).border = Border(
                    left=BORDE_FINO if columna == inicio else Side(),
                    right=BORDE_FINO if columna == fin else Side(),
                    top=BORDE_FINO if fila == 12 else Side(),
                    bottom=BORDE_FINO if fila == 14 else Side(),
                )


def escribir_resumen(hoja, inventario: Inventario, fisico_por_ubicacion: dict) -> None:
    hoja.title = "Resumen"
    hoja.sheet_view.showGridLines = False
    aplicar_titulo(hoja, inventario)
    escribir_datos_generales(hoja, inventario)
    escribir_indicadores(hoja, inventario)

    es_medios = inventario.modalidad_conteo == "MEDIOS"
    encabezados = (
        [
            "Cód. CGR",
            "Cód. cliente",
            "Descripción",
            "Stock sistema",
            "Vacíos",
            "Llenos",
            "Rotos",
            "Utilizables",
            "Total físico",
            "Diferencia",
            "Físico / teórico",
            "Resultado",
            "Ajustado sistema",
            "Fecha ajuste sistema",
            "Causa",
            "Observación",
        ]
        if es_medios
        else [
            "Cód. CGR",
            "Cód. cliente",
            "Descripción",
            "Etapa",
            "Operación / lote",
            "Proveedor",
            "Stock sistema",
            "Disponible",
            "En línea",
            "Scrap",
            "Sin ubicación / ajuste PC",
            "Diferencia",
            "Físico / teórico",
            "Resultado",
            "Ajustado sistema",
            "Fecha ajuste sistema",
            "Causa",
            "Observación",
        ]
    )
    fila_encabezado = 16
    hoja.append([])
    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(fila_encabezado, columna, encabezado)
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.font = Font(
            name="Arial",
            size=8,
            bold=True,
            color=BLANCO,
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        celda.border = Border(
            left=BORDE_AZUL,
            right=BORDE_AZUL,
            top=BORDE_AZUL,
            bottom=BORDE_AZUL,
        )
    hoja.row_dimensions[fila_encabezado].height = 26

    for fila, item in enumerate(
        inventario.items,
        start=fila_encabezado + 1,
    ):
        metricas = metricas_conteo_item(item)
        ubicaciones = fisico_por_ubicacion.get(item.id, {})
        relacion = (
            metricas["relacion_porcentaje"] / 100
            if metricas["relacion_porcentaje"] is not None
            else None
        )
        valores = (
            [
                item.codigo,
                item.codigo_cliente,
                item.descripcion,
                item.stock_sistema,
                item.conteo_vacios,
                item.conteo_llenos,
                item.conteo_rotos,
                item.conteo_vacios,
                item.conteo_fisico,
                metricas["diferencia"],
                relacion,
                resultado_item(item, metricas),
                item.ajustado_sistema,
                item.ajustado_sistema_en,
                item.causa_diferencia_codigo or "",
                item.observaciones or "",
            ]
            if es_medios
            else [
                codigo_resumen(item),
                "" if item.etapa == "MATERIAL" else item.codigo_cliente,
                descripcion_resumen(item),
                texto_etapa(item.etapa),
                texto_operacion(item),
                item.proveedor or "",
                (
                    "Pendiente de integrar"
                    if item.stock_pendiente
                    else item.stock_sistema
                ),
                ubicaciones.get("DISPONIBLE"),
                ubicaciones.get("LINEA"),
                ubicaciones.get("SCRAP"),
                ubicaciones.get("SIN_UBICACION"),
                metricas["diferencia"],
                relacion,
                resultado_item(item, metricas),
                item.ajustado_sistema,
                item.ajustado_sistema_en,
                item.causa_diferencia_codigo or "",
                item.observaciones or "",
            ]
        )
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(fila, columna, valor)
            celda.font = Font(name="Arial", size=8, color="101828")
            celda.alignment = Alignment(
                horizontal=(
                    "center"
                    if encabezados[columna - 1] in {
                        "Cód. CGR", "Cód. cliente", "Etapa",
                        "Operación / lote", "Proveedor", "Resultado",
                    }
                    else "right"
                    if encabezados[columna - 1] in {
                        "Stock sistema", "Disponible", "En línea", "Scrap",
                        "Sin ubicación / ajuste PC", "Vacíos",
                        "Llenos", "Rotos", "Utilizables", "Total físico",
                        "Diferencia", "Físico / teórico",
                    }
                    else "left"
                ),
                vertical="center",
                wrap_text=encabezados[columna - 1] in {
                    "Descripción", "Operación / lote", "Causa", "Observación",
                },
            )
            celda.border = Border(
                left=BORDE_FINO,
                right=BORDE_FINO,
                top=BORDE_FINO,
                bottom=BORDE_FINO,
            )
        hoja.cell(fila, 1).font = Font(
            name="Arial",
            size=8,
            bold=True,
            color=AZUL_OSCURO,
        )
        if not es_medios and item.etapa == "MATERIAL":
            columnas_numericas = {
                encabezados.index(nombre) + 1
                for nombre in {
                    "Stock sistema", "Disponible", "En línea", "Scrap",
                    "Sin ubicación / ajuste PC", "Diferencia",
                    "Físico / teórico",
                }
            }
            for columna in range(1, len(encabezados) + 1):
                if columna not in columnas_numericas:
                    hoja.cell(fila, columna).fill = PatternFill(
                        "solid",
                        fgColor=GRIS,
                    )
        for nombre in {
            "Stock sistema", "Disponible", "En línea", "Scrap",
            "Sin ubicación / ajuste PC", "Vacíos", "Llenos",
            "Rotos", "Utilizables", "Total físico", "Diferencia",
        }:
            if nombre in encabezados:
                hoja.cell(
                    fila, encabezados.index(nombre) + 1
                ).number_format = "#,##0.00"
        hoja.cell(
            fila, encabezados.index("Físico / teórico") + 1
        ).number_format = "0.00%"
        hoja.cell(
            fila, encabezados.index("Fecha ajuste sistema") + 1
        ).number_format = "dd/mm/yyyy hh:mm"

        resultado = hoja.cell(fila, encabezados.index("Resultado") + 1)
        if resultado.value == "Dentro de ±5%":
            resultado.fill = PatternFill("solid", fgColor=VERDE)
            resultado.font = Font(
                name="Arial",
                size=8,
                bold=True,
                color=VERDE_TEXTO,
            )
        elif resultado.value == "Fuera de ±5%":
            resultado.fill = PatternFill("solid", fgColor=ROJO)
            resultado.font = Font(
                name="Arial",
                size=8,
                bold=True,
                color=ROJO_TEXTO,
            )
        else:
            resultado.fill = PatternFill("solid", fgColor=GRIS)
            resultado.font = Font(
                name="Arial",
                size=8,
                bold=True,
                color=GRIS_TEXTO,
            )
        hoja.row_dimensions[fila].height = 30

    ultima_fila = fila_encabezado + len(inventario.items)
    if inventario.items:
        tabla = Table(
            displayName="ResumenConteo",
            ref=(
                f"A{fila_encabezado}:"
                f"{get_column_letter(len(encabezados))}{ultima_fila}"
            ),
        )
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        hoja.add_table(tabla)

    anchos = (
        [15, 16, 29, 14, 12, 12, 12, 14, 14, 13, 16, 18, 20, 20, 22, 36]
        if es_medios
        else [
            15, 16, 31, 18, 25, 18, 15, 13, 13,
            13, 20, 13, 16, 18, 20, 20, 22, 38,
        ]
    )
    for columna, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho

    hoja.freeze_panes = "A17"
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = f"{fila_encabezado}:{fila_encabezado}"
    hoja.print_area = (
        f"A1:{get_column_letter(len(encabezados))}{ultima_fila}"
    )
    hoja.sheet_view.zoomScale = 85


def columnas_detalle() -> list[str]:
    return [
        "Inventario ID",
        "Código inventario",
        "Origen inventario",
        "Plantilla ID",
        "Programación ID",
        "Cliente",
        "Fecha programada",
        "Fecha realizada",
        "Mes de cierre",
        "Fecha cierre",
        "Responsable",
        "Colaboradores",
        "Ámbito",
        "Estado",
        "Categoría",
        "Modalidad de conteo",
        "Controla remitos",
        "Último remito ingreso",
        "Último remito egreso",
        "Comentario inventario",
        "Stock actualizado",
        "Fuente stock",
        "Confiabilidad inventario",
        "Renglones OK",
        "Renglones evaluables",
        "Ítem ID",
        "Código pieza",
        "Código cliente",
        "Código material",
        "Descripción",
        "Ubicación",
        "Diámetro",
        "Largo",
        "Espesor",
        "Unidad",
        "Lote",
        "Cuenta material",
        "Código padre",
        "Nivel estructura",
        "Uso acumulado",
        "Ruta estructura",
        "Tipo estructura",
        "Etapa",
        "Operación",
        "Nombre operación",
        "Proveedor",
        "Stock pendiente",
        "Stock sistema",
        "Conteo físico",
        "Físico disponible",
        "Físico en línea",
        "Físico scrap",
        "Físico sin ubicación / ajuste PC",
        "Vacíos",
        "Llenos",
        "Rotos",
        "Medios utilizables",
        "Diferencia",
        "Físico / teórico",
        "Desvío",
        "Dentro de ±5%",
        "Ajustado sistema",
        "Fecha ajuste sistema",
        "Código causa diferencia",
        "Observación del renglón",
        "Exportado el",
    ]


def fila_detalle(
    inventario: Inventario,
    item,
    confianza: dict,
    exportado_en: datetime,
    fisico_por_ubicacion: dict,
) -> list:
    metricas = metricas_conteo_item(item)
    ubicaciones = fisico_por_ubicacion.get(item.id, {})
    return [
        inventario.id,
        inventario.codigo,
        getattr(inventario, "origen", "PLANTILLA"),
        inventario.plantilla_id,
        inventario.programacion_id,
        inventario.cliente,
        inventario.fecha_programada,
        inventario.fecha,
        inventario.mes_cierre,
        inventario.cerrado_en,
        inventario.responsable,
        inventario.colaboradores or "",
        inventario.ambito,
        inventario.estado,
        inventario.categoria,
        inventario.modalidad_conteo,
        inventario.controla_remitos,
        inventario.ultimo_remito_ingreso or "",
        inventario.ultimo_remito_egreso or "",
        inventario.observaciones or "",
        inventario.stock_actualizado_en,
        inventario.stock_origen,
        (
            confianza["confiabilidad"] / 100
            if confianza["confiabilidad"] is not None
            else None
        ),
        confianza["items_ok"],
        confianza["items_evaluables"],
        item.id,
        item.codigo,
        item.codigo_cliente,
        item.codigo_material,
        item.descripcion,
        item.ubicacion,
        item.diametro,
        item.largo,
        item.espesor,
        item.unidad,
        item.lote,
        item.contar_material,
        item.codigo_padre,
        item.nivel,
        item.uso,
        item.ruta,
        item.tipo_estructura,
        item.etapa,
        item.operacion,
        item.titulo_operacion,
        item.proveedor,
        item.stock_pendiente,
        item.stock_sistema,
        item.conteo_fisico,
        ubicaciones.get("DISPONIBLE"),
        ubicaciones.get("LINEA"),
        ubicaciones.get("SCRAP"),
        ubicaciones.get("SIN_UBICACION"),
        item.conteo_vacios,
        item.conteo_llenos,
        item.conteo_rotos,
        item.conteo_vacios,
        metricas["diferencia"],
        (
            metricas["relacion_porcentaje"] / 100
            if metricas["relacion_porcentaje"] is not None
            else None
        ),
        (
            metricas["desvio_porcentaje"] / 100
            if metricas["desvio_porcentaje"] is not None
            else None
        ),
        metricas["dentro_tolerancia"],
        item.ajustado_sistema,
        item.ajustado_sistema_en,
        item.causa_diferencia_codigo or "",
        item.observaciones or "",
        exportado_en,
    ]


def escribir_detalle_historico(
    hoja,
    inventario: Inventario,
    fisico_por_ubicacion: dict,
) -> None:
    hoja.title = "Detalle histórico"
    hoja.sheet_view.showGridLines = False
    encabezados = columnas_detalle()
    confianza = metricas_confiabilidad(inventario)
    exportado_en = datetime.now()

    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(1, columna, encabezado)
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.font = Font(
            name="Arial",
            size=8,
            bold=True,
            color=BLANCO,
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        celda.border = Border(
            left=BORDE_AZUL,
            right=BORDE_AZUL,
            top=BORDE_AZUL,
            bottom=BORDE_AZUL,
        )
    hoja.row_dimensions[1].height = 32

    for fila, item in enumerate(inventario.items, start=2):
        for columna, valor in enumerate(
            fila_detalle(
                inventario,
                item,
                confianza,
                exportado_en,
                fisico_por_ubicacion,
            ),
            start=1,
        ):
            celda = hoja.cell(fila, columna, valor)
            celda.font = Font(name="Arial", size=8, color="101828")
            celda.alignment = Alignment(
                vertical="center",
                wrap_text=encabezados[columna - 1] in {
                    "Responsable", "Colaboradores", "Último remito ingreso",
                    "Último remito egreso", "Comentario inventario",
                    "Código cliente", "Descripción", "Ruta estructura",
                    "Nombre operación", "Código causa diferencia",
                    "Observación del renglón",
                },
            )
        hoja.row_dimensions[fila].height = 24

    ultima_fila = 1 + len(inventario.items)
    if inventario.items:
        tabla = Table(
            displayName="DetalleHistoricoInventario",
            ref=(
                f"A1:{get_column_letter(len(encabezados))}{ultima_fila}"
            ),
        )
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        hoja.add_table(tabla)

    columnas_por_nombre = {
        nombre: indice
        for indice, nombre in enumerate(encabezados, start=1)
    }
    for nombre in ["Fecha programada", "Fecha realizada"]:
        columna = columnas_por_nombre[nombre]
        for fila in range(2, ultima_fila + 1):
            hoja.cell(fila, columna).number_format = "dd/mm/yyyy"
    columna_mes = columnas_por_nombre["Mes de cierre"]
    for fila in range(2, ultima_fila + 1):
        hoja.cell(fila, columna_mes).number_format = "mmmm yyyy"
    for nombre in [
        "Fecha cierre",
        "Stock actualizado",
        "Fecha ajuste sistema",
        "Exportado el",
    ]:
        columna = columnas_por_nombre[nombre]
        for fila in range(2, ultima_fila + 1):
            hoja.cell(fila, columna).number_format = "dd/mm/yyyy hh:mm"
    for nombre in [
        "Confiabilidad inventario",
        "Físico / teórico",
        "Desvío",
    ]:
        columna = columnas_por_nombre[nombre]
        for fila in range(2, ultima_fila + 1):
            hoja.cell(fila, columna).number_format = "0.00%"
    for nombre in [
        "Uso acumulado",
        "Stock sistema",
        "Conteo físico",
        "Físico disponible",
        "Físico en línea",
        "Físico scrap",
        "Físico sin ubicación / ajuste PC",
        "Vacíos",
        "Llenos",
        "Rotos",
        "Medios utilizables",
        "Diferencia",
    ]:
        columna = columnas_por_nombre[nombre]
        for fila in range(2, ultima_fila + 1):
            hoja.cell(fila, columna).number_format = "#,##0.00"

    anchos_personalizados = {
        "Inventario ID": 13,
        "Código inventario": 18,
        "Origen inventario": 17,
        "Cliente": 22,
        "Fecha cierre": 22,
        "Responsable": 28,
        "Modalidad de conteo": 22,
        "Controla remitos": 22,
        "Último remito ingreso": 24,
        "Último remito egreso": 24,
        "Comentario inventario": 42,
        "Renglones evaluables": 19,
        "Ítem ID": 17,
        "Código pieza": 18,
        "Código cliente": 24,
        "Código material": 18,
        "Descripción": 34,
        "Lote": 18,
        "Ruta estructura": 42,
        "Tipo estructura": 18,
        "Nombre operación": 28,
        "Físico sin ubicación / ajuste PC": 25,
        "Código causa diferencia": 25,
        "Observación del renglón": 44,
    }
    for columna, encabezado in enumerate(encabezados, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = (
            anchos_personalizados.get(encabezado, 14)
        )

    hoja.freeze_panes = "F2"
    hoja.sheet_view.zoomScale = 75


def generar_excel(inventario: Inventario, cajas: list[dict] | None = None) -> BytesIO:
    fisico_por_ubicacion = desglosar_conteo_por_ubicacion(
        inventario.items,
        cajas or [],
    )
    libro = Workbook()
    resumen = libro.active
    detalle = libro.create_sheet("Detalle histórico")
    escribir_resumen(resumen, inventario, fisico_por_ubicacion)
    escribir_detalle_historico(detalle, inventario, fisico_por_ubicacion)
    libro.active = 0

    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    return archivo
