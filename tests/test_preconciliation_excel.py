from datetime import date
from types import SimpleNamespace
from pathlib import Path
from uuid import uuid4
import unittest

from openpyxl import load_workbook

from app.services.preconciliation_excel import crear_excel_preconciliacion
from app import preconciliation_audit as audit
from test_preconciliation_audit import calculate


class PreconciliationExcelTests(unittest.TestCase):
    def setUp(self):
        self.old_path = audit.STORE_PATH
        self.audit_path = Path(__file__).resolve().parent / f"{uuid4().hex}.sqlite3"
        audit.STORE_PATH = self.audit_path
        self.final = SimpleNamespace(codigo="INV-FINAL", cliente="Cliente", fecha=date(2026, 8, 31))
        self.calculation = calculate()
        for scenario in self.calculation["escenarios"]:
            for row in scenario["filas"]:
                row["codigo_cliente"] = f"CLI-{row['codigo']}"
                for detail in row.get("detalle_ingresos", []):
                    detail["codigo_cliente_origen"] = ""
                for detail in row.get("detalle_egresos", []):
                    detail["codigo_cliente_origen"] = "CLI-PADRE"
        for movement in self.calculation["movimientos"]:
            movement["codigo_cliente"] = (
                "CLI-PADRE" if movement["tipo"] == "OUT" else ""
            )
        self.notes = {"P1": {"estado": "REVISAR_MOVIMIENTOS", "causa": "MOV", "observacion": "Controlar remito"}}
        self.history = [{"codigo": "INV-ANT", "fecha": "2026-05-31", "filas": {"P1": {"conteo": 118, "sistema": 120, "diferencia": -2}}}]

    def tearDown(self):
        audit.STORE_PATH = self.old_path
        self.audit_path.unlink(missing_ok=True)

    def test_analysis_export_has_formulas_and_all_sections(self):
        book = load_workbook(crear_excel_preconciliacion(self.final, self.calculation, self.notes, self.history, "ANALISIS"), data_only=False)
        sheet = book["Análisis general"]
        headers = [cell.value for cell in sheet[4]]
        self.assertIn("Sistema MP equivalente", headers)
        self.assertIn("Físico MP equivalente", headers)
        self.assertIn("INV-ANT diferencia", headers)
        self.assertIn("INV-JUN esperado explotado", headers)
        self.assertIn("INV-JUN físico explotado vs esperado", headers)
        self.assertNotIn("INV-JUN sistema vs esperado", headers)
        self.assertEqual(headers[:3], ["Código interno", "Código cliente", "Descripción"])
        self.assertEqual(sheet["B5"].value, "CLI-P1")
        self.assertEqual(sheet["R5"].value, "=M5-E5")
        self.assertIn("Indicadores", book.sheetnames)
        self.assertEqual(book["Indicadores"]["A5"].value, "Confiabilidad vs. sistema")

    def test_reconciliation_export_contains_full_detail(self):
        calculation = dict(self.calculation, escenarios=[self.calculation["escenarios"][0]])
        book = load_workbook(crear_excel_preconciliacion(self.final, calculation, self.notes, self.history, "CONCILIACION"), data_only=False)
        self.assertEqual(book.sheetnames, ["Conciliación", "Ingresos", "Egresos", "Auditoría remitos", "Ajustes sistema", "Indicadores"])
        self.assertEqual(book["Conciliación"]["H5"].value, "=D5+E5-F5-G5")
        self.assertEqual(book["Conciliación"]["J5"].value, "=I5-H5")
        headers = [cell.value for cell in book["Conciliación"][4]]
        self.assertNotIn("Sistema vs esperado", headers)
        self.assertEqual(headers[:3], ["Código interno", "Código cliente", "Descripción"])
        self.assertEqual(book["Conciliación"]["B5"].value, "CLI-P1")
        self.assertGreater(book["Ingresos"].max_row, 1)
        self.assertGreater(book["Egresos"].max_row, 1)
        self.assertEqual(book["Ingresos"]["B2"].value, "CLI-P1")
        self.assertEqual(book["Egresos"]["B2"].value, "CLI-P1")
        self.assertEqual(book["Egresos"]["G2"].value, "CLI-PADRE")
        self.assertEqual(book["Auditoría remitos"]["E3"].value, "CLI-PADRE")
        self.assertEqual(book["Ajustes sistema"]["B1"].value, "Código cliente")


if __name__ == "__main__":
    unittest.main()
